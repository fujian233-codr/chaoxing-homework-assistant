from __future__ import annotations

import re
from pathlib import Path

from docx import Document

from .models import Assignment
from .storage import Storage
from .utils import safe_read_text


def markdown_text_to_paragraphs(text: str) -> list[tuple[str, str]]:
    parsed: list[tuple[str, str]] = []
    in_code_block = False
    code_lines: list[str] = []
    for raw_line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        line = raw_line.rstrip()
        if line.strip().startswith("```"):
            if in_code_block:
                parsed.append(("code", "\n".join(code_lines)))
                code_lines = []
            in_code_block = not in_code_block
            continue
        if in_code_block:
            code_lines.append(raw_line)
            continue
        if not line.strip():
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", line)
        if heading:
            parsed.append((f"heading-{len(heading.group(1))}", heading.group(2).strip()))
        elif re.match(r"^[-*+]\s+", line):
            parsed.append(("bullet", re.sub(r"^[-*+]\s+", "", line).strip()))
        else:
            parsed.append(("paragraph", line.strip()))
    if code_lines:
        parsed.append(("code", "\n".join(code_lines)))
    return parsed


def build_docx_from_draft(
    assignment: Assignment,
    draft_path: Path,
    output_path: Path,
    force: bool = False,
) -> Path:
    if output_path.exists() and output_path.stat().st_size > 0 and not force:
        raise FileExistsError(f"目标 Word 已存在且非空：{output_path}")
    if not draft_path.is_file():
        raise FileNotFoundError(f"草稿文件不存在：{draft_path}")
    draft = safe_read_text(draft_path).strip()
    if not draft:
        raise ValueError("草稿内容为空")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading(assignment.title, level=0)
    document.add_paragraph(f"课程：{assignment.course_name}")
    if assignment.requirements:
        requirement = document.add_paragraph()
        requirement.add_run("作业要求摘录：").bold = True
        requirement.add_run(assignment.requirements[:1200])

    parsed = markdown_text_to_paragraphs(draft)
    for kind, content in parsed:
        if kind == "heading-1":
            document.add_heading(content, level=1)
        elif kind in {"heading-2", "heading-3", "heading-4"}:
            document.add_heading(content, level=2)
        elif kind == "bullet":
            document.add_paragraph(content, style="List Bullet")
        elif kind == "code":
            paragraph = document.add_paragraph()
            run = paragraph.add_run(content)
            run.font.name = "Consolas"
            run.font.size = None
        else:
            document.add_paragraph(content)
    document.save(output_path)
    return output_path


def build_docx_command(
    storage: Storage,
    account_id: str,
    assignment_id: str,
    draft_path: Path,
    output_path: Path | None,
    force: bool,
) -> Path:
    assignment = storage.get_assignment(account_id, assignment_id)
    if not assignment:
        raise KeyError(f"数据库中未找到作业：账号 {account_id} / 作业 {assignment_id}")
    target = output_path or Path(assignment.output_path)
    if not target:
        raise ValueError("未指定输出路径，且数据库没有记录输出路径")
    result = build_docx_from_draft(assignment, draft_path, target, force=force)
    storage.update_assignment_output(account_id, assignment_id, str(result.resolve()))
    storage.record_ai_task(
        account_id=account_id,
        assignment_id=assignment_id,
        status="draft_generated",
        output_path=str(result.resolve()),
        draft_path=str(draft_path.resolve()),
    )
    return result
